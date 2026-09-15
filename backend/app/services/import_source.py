"""Bounded CSV/XLSX iteration; only the compatibility parser builds a list."""
import csv
import io
from contextlib import contextmanager
from datetime import date, datetime
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook

MAX_ROWS = 20_000
MAX_COLUMNS = 100
MAX_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024


@contextmanager
def source_rows(source, extension):
    workbook = None
    stream = None
    try:
        if extension == ".csv":
            stream = io.TextIOWrapper(source, encoding="utf-8-sig", newline="")
            sample = stream.read(8192)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = csv.reader(stream, dialect)
        elif extension == ".xlsx":
            with ZipFile(source) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                    raise HTTPException(413, "Expanded workbook exceeds 100 MB.")
            source.seek(0)
            workbook = load_workbook(source, read_only=True, data_only=False)
            sheet = workbook.active
            if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                raise HTTPException(422, "Import files cannot contain more than 100 columns.")
            # Bound iteration independently of untrusted worksheet dimension metadata.
            sheet.reset_dimensions()
            rows = sheet.iter_rows()
        else:
            raise HTTPException(415, "Only .csv and .xlsx import files are supported.")

        def values(row):
            if extension == ".xlsx":
                if any(cell.data_type == "f" for cell in row):
                    raise HTTPException(422, "Import formulas are not supported; export values first.")
                row = [cell.value for cell in row]
            if len(row) > MAX_COLUMNS:
                raise HTTPException(422, "Import files cannot contain more than 100 columns.")
            return row

        from app.services.imports import normalized_header
        first = next(rows, None)
        if first is None:
            raise HTTPException(422, "The file must contain a header row.")
        headers = [str(v or "").strip() for v in values(first)]
        if not headers or any(not h for h in headers):
            raise HTTPException(422, "Every source column must have a header.")
        if len({normalized_header(h) for h in headers}) != len(headers):
            raise HTTPException(422, "Source column headers must be unique.")

        def records():
            for number, row in enumerate(rows, 2):
                if number > MAX_ROWS + 1:
                    raise HTTPException(413, "Import files cannot contain more than 20000 data rows.")
                row = list(values(row))
                if not any(v is not None and str(v).strip() for v in row):
                    continue
                if any(v is not None and str(v).strip() for v in row[len(headers):]):
                    raise HTTPException(422, "Data extends beyond the header columns.")
                row += [None] * (len(headers) - len(row))
                yield number, {key: value.isoformat() if isinstance(value, (date, datetime)) else "" if value is None else value
                               for key, value in zip(headers, row)}
        yield headers, records()
    except HTTPException:
        raise
    except (UnicodeError, csv.Error) as exc:
        raise HTTPException(422, "CSV files must contain valid UTF-8 CSV data.") from exc
    except Exception as exc:
        raise HTTPException(422, "The import source could not be read.") from exc
    finally:
        if workbook:
            workbook.close()
        if stream:
            stream.detach()
