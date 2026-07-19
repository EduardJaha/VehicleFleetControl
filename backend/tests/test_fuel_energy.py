import asyncio
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from starlette.datastructures import Headers, UploadFile

from app.api.v1.endpoints.fuel import (
    add_fuel_record,
    delete_fuel_record,
    overview,
    overview_totals,
    restore_fuel_record,
    update_fuel_record,
)
from app.api.v1.endpoints.reports import build_excel, fuel_costs_report
from app.core.security import require_roles
from app.db.session import Base
from app.models import Attachment, AuditLog, User, Vehicle, VehicleFuel
from app.schemas import EnergyUnit, FuelUpdate, UserRole
from app.utils import files as file_utils


@pytest.fixture()
def db() -> Session:
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = Session(engine)
    yield session
    session.close()
    engine.dispose()


@pytest.fixture()
def users(db: Session) -> dict[str, User]:
    rows = {
        role: User(
            email=f"{role}@example.com",
            full_name=role.title(),
            hashed_password="unused",
            role=role,
            is_active=True,
        )
        for role in ("admin", "finance", "viewer")
    }
    db.add_all(rows.values())
    db.commit()
    return rows


def make_vehicle(
    db: Session,
    fuel_type: str,
    *,
    sequence: int | None = None,
    odometer: int = 10_000,
    archived: bool = False,
) -> Vehicle:
    sequence = sequence or (200 + db.query(Vehicle).count())
    vehicle = Vehicle(
        brand="Test",
        model=f"{fuel_type} Model",
        fuel_type=fuel_type,
        vehicle_location="Prishtina",
        license_plate=f"01-{sequence:03d}-AB",
        registration_country="XK",
        license_plate_normalized=f"01{sequence:03d}AB",
        status=0,
        odometer_km=odometer,
        archived=archived,
    )
    db.add(vehicle)
    db.commit()
    return vehicle


def create_record(
    db: Session,
    user: User,
    vehicle: Vehicle,
    *,
    quantity: Decimal | None = Decimal("45.500"),
    unit_cost: Decimal | None = Decimal("1.3500"),
    odometer: int | None = 10_100,
    submitted_fuel_type: str | None = None,
    liters: Decimal | None = None,
    cost_per_liter: Decimal | None = None,
):
    return asyncio.run(add_fuel_record(
        refuel_date="19-07-2026",
        vehicle_id=vehicle.id,
        license_plate=None,
        quantity=quantity,
        unit_cost=unit_cost,
        location=None,
        station_name="Test Station",
        odometer_km=odometer,
        bill_file=None,
        fuel_type=submitted_fuel_type,
        unit=None,
        liters=liters,
        cost_per_liter=cost_per_liter,
        db=db,
        current_user=user,
    ))


@pytest.mark.parametrize("fuel_type", ["Petrol", "Diesel", "Hybrid", "LPG", "CNG", "Gas"])
def test_liquid_vehicle_type_is_authoritative_and_uses_liters(
    db: Session,
    users,
    fuel_type: str,
):
    vehicle = make_vehicle(db, fuel_type)
    result = create_record(db, users["finance"], vehicle)
    record = result["record"]

    assert record.fuel_type == fuel_type
    assert record.unit == EnergyUnit.liter
    assert record.quantity == Decimal("45.500")
    stored = db.get(VehicleFuel, record.id)
    assert stored.liters is None
    assert stored.cost_per_liter is None


def test_electric_vehicle_derives_kwh_and_decimal_total(db: Session, users):
    vehicle = make_vehicle(db, "Electric")
    result = create_record(
        db,
        users["admin"],
        vehicle,
        quantity=Decimal("62.400"),
        unit_cost=Decimal("0.1799"),
    )
    record = result["record"]

    assert record.fuel_type == "Electric"
    assert record.unit == EnergyUnit.kilowatt_hour
    assert record.unit_cost == Decimal("0.1799")
    assert record.total_cost == Decimal("11.23")
    stored = db.get(VehicleFuel, record.id)
    assert stored.liters is None
    assert stored.cost_per_liter is None


def test_frontend_does_not_need_to_submit_fuel_type_and_mismatch_is_rejected(db: Session, users):
    diesel = make_vehicle(db, "Diesel")
    assert create_record(db, users["finance"], diesel)["record"].fuel_type == "Diesel"

    electric = make_vehicle(db, "Electric")
    with pytest.raises(HTTPException) as mismatch:
        create_record(
            db,
            users["finance"],
            electric,
            submitted_fuel_type="Diesel",
        )
    assert mismatch.value.status_code == 400
    assert "configured as Electric" in mismatch.value.detail


def test_archived_vehicle_cannot_receive_record(db: Session, users):
    vehicle = make_vehicle(db, "Diesel", archived=True)
    with pytest.raises(HTTPException) as archived:
        create_record(db, users["finance"], vehicle)
    assert archived.value.status_code == 400
    assert "archived Vehicle" in archived.value.detail


@pytest.mark.parametrize("quantity", [Decimal("0"), Decimal("-0.001")])
def test_nonpositive_quantity_is_rejected_with_unit_aware_message(
    db: Session,
    users,
    quantity: Decimal,
):
    electric = make_vehicle(db, "Electric")
    with pytest.raises(HTTPException) as invalid:
        create_record(db, users["admin"], electric, quantity=quantity)
    assert invalid.value.detail == "Charging energy in kWh must be greater than zero."


def test_negative_unit_cost_and_electric_legacy_liters_are_rejected(db: Session, users):
    diesel = make_vehicle(db, "Diesel")
    with pytest.raises(HTTPException) as negative:
        create_record(
            db,
            users["finance"],
            diesel,
            unit_cost=Decimal("-0.0001"),
        )
    assert negative.value.detail == "Cost per liter cannot be negative."

    electric = make_vehicle(db, "Electric")
    with pytest.raises(HTTPException) as legacy:
        create_record(
            db,
            users["finance"],
            electric,
            quantity=None,
            unit_cost=None,
            liters=Decimal("12.5"),
            cost_per_liter=Decimal("0.2"),
        )
    assert "must use quantity in kWh" in legacy.value.detail


def test_odometer_equal_or_higher_is_accepted_and_lower_is_rejected(db: Session, users):
    equal_vehicle = make_vehicle(db, "Diesel", odometer=10_000)
    create_record(db, users["finance"], equal_vehicle, odometer=10_000)
    assert db.get(Vehicle, equal_vehicle.id).odometer_km == 10_000

    higher_vehicle = make_vehicle(db, "Electric", odometer=20_000)
    create_record(db, users["finance"], higher_vehicle, odometer=20_250)
    assert db.get(Vehicle, higher_vehicle.id).odometer_km == 20_250

    lower_vehicle = make_vehicle(db, "Petrol", odometer=30_000)
    with pytest.raises(HTTPException) as lower:
        create_record(db, users["finance"], lower_vehicle, odometer=29_999)
    assert lower.value.status_code == 400
    assert "30,000 km" in lower.value.detail


def test_edit_preserves_historical_type_and_unit_and_recalculates_total(db: Session, users):
    vehicle = make_vehicle(db, "Electric")
    created = create_record(
        db,
        users["admin"],
        vehicle,
        quantity=Decimal("10.000"),
        unit_cost=Decimal("0.2000"),
    )["record"]
    vehicle.fuel_type = "Diesel"
    db.commit()

    update_fuel_record(
        created.id,
        FuelUpdate(
            refuel_date="20-07-2026",
            quantity=Decimal("12.500"),
            unit_cost=Decimal("0.3333"),
            location="Tirana",
            station_name="Provider",
            odometer_km=10_200,
        ),
        db,
        users["admin"],
    )
    record = db.get(VehicleFuel, created.id)
    assert record.fuel_type == "Electric"
    assert record.unit == "KWH"
    assert record.quantity == Decimal("12.500")
    assert record.total_cost == Decimal("4.17")

    with pytest.raises(HTTPException):
        update_fuel_record(
            created.id,
            FuelUpdate(
                refuel_date="20-07-2026",
                quantity=Decimal("12.500"),
                fuel_type="Diesel",
            ),
            db,
            users["admin"],
        )


def test_overview_and_reports_never_add_liters_and_kwh(db: Session, users):
    diesel = make_vehicle(db, "Diesel")
    electric = make_vehicle(db, "Electric")
    create_record(
        db,
        users["finance"],
        diesel,
        quantity=Decimal("40.000"),
        unit_cost=Decimal("1.5000"),
    )
    create_record(
        db,
        users["finance"],
        electric,
        quantity=Decimal("60.000"),
        unit_cost=Decimal("0.2000"),
    )

    groups = overview(db=db)
    assert {(group.unit, group.total_quantity) for group in groups} == {
        (EnergyUnit.liter, Decimal("40.000")),
        (EnergyUnit.kilowatt_hour, Decimal("60.000")),
    }
    totals = overview_totals(db=db)
    assert totals.total_liters == Decimal("40.000")
    assert totals.total_kwh == Decimal("60.000")
    assert totals.total_fuel_cost == Decimal("72.00")

    report = fuel_costs_report(db)
    assert report["kpis"]["total_liters"] == 40.0
    assert report["kpis"]["total_kwh"] == 60.0
    assert report["kpis"]["total_fuel_cost"] == 72.0
    assert report["kpis"]["total_fuel_and_charging_cost"] == 72.0
    assert {row["unit"] for row in report["rows"]} == {"L", "KWH"}


def test_excel_export_uses_generic_numeric_columns(db: Session, users):
    vehicle = make_vehicle(db, "Electric")
    create_record(db, users["finance"], vehicle)
    report = fuel_costs_report(db)
    workbook = load_workbook(BytesIO(build_excel("fuel-costs", report).getvalue()))
    sheet = workbook["Rows"]
    headers = [cell.value for cell in sheet[1]]

    assert "Quantity" in headers
    assert "Unit" in headers
    assert "Unit Price" in headers
    assert "Liters" not in headers
    assert isinstance(sheet.cell(2, headers.index("Quantity") + 1).value, (int, float))


def test_audit_snapshot_is_unit_aware_and_permissions_remain_intact(db: Session, users):
    vehicle = make_vehicle(db, "Electric")
    result = create_record(db, users["finance"], vehicle)
    log = db.query(AuditLog).filter(
        AuditLog.entity_type == "VehicleFuel",
        AuditLog.entity_id == result["record"].id,
    ).one()
    assert log.new_values["quantity"] == "45.500"
    assert log.new_values["unit"] == "KWH"
    assert "liters" not in log.new_values
    assert "cost_per_liter" not in log.new_values
    assert "bill_file_path" not in log.new_values

    dependency = require_roles(UserRole.admin, UserRole.finance)
    assert dependency(users["admin"]).role == "admin"
    assert dependency(users["finance"]).role == "finance"
    with pytest.raises(HTTPException) as denied:
        dependency(users["viewer"])
    assert denied.value.status_code == 403


def test_bill_upload_archive_and_restore_remain_functional(
    db: Session,
    users,
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(file_utils.settings, "upload_directory", str(tmp_path))
    vehicle = make_vehicle(db, "Electric")
    bill = UploadFile(
        file=BytesIO(b"%PDF-1.7\ncharging receipt"),
        filename="charging-receipt.pdf",
        headers=Headers({"content-type": "application/pdf"}),
    )
    result = asyncio.run(add_fuel_record(
        refuel_date="19-07-2026",
        vehicle_id=vehicle.id,
        license_plate=None,
        quantity=Decimal("20"),
        unit_cost=Decimal("0.25"),
        location="Tirana",
        station_name="Public Charger",
        odometer_km=10_100,
        bill_file=bill,
        fuel_type=None,
        unit=None,
        liters=None,
        cost_per_liter=None,
        db=db,
        current_user=users["finance"],
    ))
    record_id = result["record"].id
    attachment = db.query(Attachment).filter(
        Attachment.entity_type == "VehicleFuel",
        Attachment.entity_id == record_id,
    ).one()
    assert result["record"].bill_file_path == f"/api/v1/files/{attachment.id}/download"

    delete_fuel_record(record_id, db, users["finance"])
    assert db.get(VehicleFuel, record_id).archived is True
    restored = restore_fuel_record(record_id, db, users["admin"])
    assert restored.archived is False
