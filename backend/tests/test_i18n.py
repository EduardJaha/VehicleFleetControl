from io import BytesIO

from openpyxl import load_workbook

from app.api.v1.endpoints.reports import build_excel
from app.core.errors import infer_error_code
from app.core.i18n import resolve_language, translate
from app.services.notifications import notification_key


def test_accept_language_resolution_supports_variants_and_fallback():
    assert resolve_language("sq-AL,sq;q=0.9,en;q=0.8") == "sq"
    assert resolve_language("en-GB") == "en"
    assert resolve_language("de-DE") == "en"
    assert resolve_language(None) == "en"


def test_error_codes_have_albanian_messages():
    assert infer_error_code(404, "/api/v1/vehicles/42") == "vehicle_not_found"
    assert translate("vehicle_not_found", "sq") == "Automjeti nuk u gjet."
    assert translate("validation_error", "sq") != translate("validation_error", "en")


def test_notification_keys_are_stable_and_language_neutral():
    assert notification_key("Work Order waiting for parts", "title") == (
        "modules:notificationContent.work_order_waiting_for_parts.title"
    )


def test_albanian_excel_localizes_labels_and_preserves_numbers():
    report = {
        "kpis": {"total_vehicles": 1},
        "rows": [{"license_plate": "01-123-AB", "status": "Active", "total_cost": 12.5}],
    }
    workbook = load_workbook(BytesIO(build_excel("fleet-summary", report, "sq").getvalue()))
    assert "Treguesit" in workbook.sheetnames
    assert "Rreshtat" in workbook.sheetnames
    sheet = workbook["Rreshtat"]
    assert [cell.value for cell in sheet[1]] == ["Targa", "Statusi", "Kostoja Totale"]
    assert sheet.cell(2, 2).value == "Aktiv"
    assert isinstance(sheet.cell(2, 3).value, (int, float))
